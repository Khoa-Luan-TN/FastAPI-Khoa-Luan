# app/services/mongo_import_service.py
from __future__ import annotations

from typing import Any, Dict, Callable, Optional, List, Set, Tuple
from datetime import datetime, timezone
import json
import logging
import os
import re
import unicodedata
from bson import ObjectId
from openpyxl import load_workbook
from app.services.keyword_alias_service import ensure_keyword_alias_indexes
from app.services.minio_marker_service import ensure_asset_prefix_markers, ensure_root_folders
from app.services.keyword_alias_service import (
    _resolve_keyword_slug,
    enforce_canonical_name_precedence,
)

_log = logging.getLogger(__name__)


IMPORT_ORDER = ["class", "subject", "topic", "lesson", "chunk", "keyword"]

REF_MAP = {
    "subject": ("class_ref", "class_id", "class"),
    "topic": ("subject_ref", "subject_id", "subject"),
    "lesson": ("topic_ref", "topic_id", "topic"),
    "chunk": ("lesson_ref", "lesson_id", "lesson"),
}

JSON_FIELDS = {"images", "videos", "image_url", "video_url"}


def _now():
    return datetime.now(timezone.utc)


def _get_import_minio():
    """Return (client, bucket) or (None, None) if MinIO is not configured."""
    bucket = (os.getenv("MINIO_BUCKET") or "").strip()
    if not bucket:
        return None, None
    try:
        from app.services.minio_client import get_minio_client
        return get_minio_client(), bucket
    except Exception:
        return None, None


def _norm_header(h: Any) -> str:
    return str(h or "").strip()


def _cell_to_value(v: Any):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s if s != "" else None
    return v


def _try_parse_json(v: Any):
    if v is None:
        return None
    if isinstance(v, (dict, list)):
        return v
    s = str(v).strip()
    if not s:
        return None
    if s.startswith("{") or s.startswith("["):
        try:
            return json.loads(s)
        except Exception:
            return s
    return s


def _slugify_vi(s: Any) -> str:
    s = ("" if s is None else str(s)).strip().lower()
    if not s:
        return ""
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.replace("đ", "d")
    s = re.sub(r"[^a-z0-9]+", "-", s)
    s = re.sub(r"-{2,}", "-", s).strip("-")
    return s


def _two_digit(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, str):
        t = v.strip()
        if t.isdigit():
            return f"{int(t):02d}"
        m = re.search(r"\d+", t)
        return f"{int(m.group()):02d}" if m else ""
    try:
        return f"{int(v):02d}"
    except Exception:
        return ""


def _read_sheet_rows(wb, sheet_name: str) -> List[Dict[str, Any]]:
    if sheet_name not in wb.sheetnames:
        return []

    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [_norm_header(x) for x in rows[0]]
    out = []

    for idx, r in enumerate(rows[1:], start=2):
        rec = {}
        empty = True
        for h, cell in zip(headers, r):
            if not h:
                continue
            val = _cell_to_value(cell)
            if val is not None:
                empty = False
            rec[h] = val
        if empty:
            continue
        rec["_row"] = idx
        out.append(rec)

    return out


def _ensure_import_index(db, col: str):
    try:
        db[col].create_index("import_key", unique=True)
    except Exception:
        pass


def _upsert_by_import_key(
    db,
    col: str,
    import_key: str,
    doc: Dict[str, Any],
    *,
    actor: str,
) -> Tuple[str, str]:
    now = _now()

    proj = {"_id": 1, "deleted_at": 1}
    for k in doc.keys():
        proj[k] = 1

    existing = db[col].find_one({"import_key": import_key}, proj)

    if existing:
        patch = dict(doc)

        if "is_deleted" in patch:
            is_del = patch["is_deleted"]
            if isinstance(is_del, str):
                is_del = is_del.strip().lower() in ("true", "1", "yes", "y", "on")
            is_del = bool(is_del)
            patch["is_deleted"] = is_del

            if is_del:
                patch["deleted_at"] = existing.get("deleted_at") or now
            else:
                patch["deleted_at"] = None

        IGNORE = {"_id", "created_at", "created_by", "updated_at", "updated_by"}
        same = True
        for k, v in patch.items():
            if k in IGNORE:
                continue
            if existing.get(k) != v:
                same = False
                break

        if same:
            return str(existing["_id"]), "noop"

        patch["updated_at"] = now
        patch["updated_by"] = actor
        db[col].update_one({"_id": existing["_id"]}, {"$set": patch})
        return str(existing["_id"]), "update"

    ins = dict(doc)
    ins["import_key"] = import_key
    ins.setdefault("is_deleted", False)
    ins.setdefault("deleted_at", None)

    ins["created_at"] = now
    ins["updated_at"] = now
    ins["created_by"] = actor
    ins["updated_by"] = actor

    if ins.get("is_deleted") is True:
        ins["deleted_at"] = now

    r = db[col].insert_one(ins)
    return str(r.inserted_id), "insert"


def _get_class_slug(db, class_ref: str, ctx: Dict[str, Any]) -> str:
    class_ref = (class_ref or "").strip()
    if not class_ref:
        return ""
    cached = ctx.get("class", {}).get(class_ref) or {}
    if cached.get("class_slug"):
        return cached["class_slug"]

    d = db["class"].find_one({"import_key": class_ref}, {"class_name": 1})
    if not d or not d.get("class_name"):
        return ""
    slug = _slugify_vi(d.get("class_name"))
    ctx.setdefault("class", {})[class_ref] = {"class_slug": slug}
    return slug


def _get_subject_path_info(db, subject_ref: str, ctx: Dict[str, Any]) -> dict:
    subject_ref = (subject_ref or "").strip()
    if not subject_ref:
        return {}
    cached = ctx.get("_subject_path", {}).get(subject_ref)
    if cached:
        return cached
    d = db["subject"].find_one(
        {"import_key": subject_ref},
        {"subject_name": 1, "class_ref": 1},
    )
    if not d:
        return {}
    class_ref = (d.get("class_ref") or "").strip()
    class_slug = _get_class_slug(db, class_ref, ctx)
    subject_slug = _slugify_vi(d.get("subject_name"))
    if not (class_slug and subject_slug):
        return {}
    info = {"class_slug": class_slug, "subject_slug": subject_slug}
    ctx.setdefault("_subject_path", {})[subject_ref] = info
    return info


def _get_topic_path_info(db, topic_ref: str, ctx: Dict[str, Any]) -> dict:
    topic_ref = (topic_ref or "").strip()
    if not topic_ref:
        return {}
    cached = ctx.get("_topic_path", {}).get(topic_ref)
    if cached:
        return cached
    d = db["topic"].find_one({"import_key": topic_ref}, {"subject_ref": 1, "topic_num": 1})
    if not d:
        return {}
    subj_info = _get_subject_path_info(db, (d.get("subject_ref") or "").strip(), ctx)
    if not subj_info:
        return {}
    topic_num = _two_digit(d.get("topic_num"))
    if not topic_num:
        return {}
    info = {**subj_info, "topic_num": topic_num}
    ctx.setdefault("_topic_path", {})[topic_ref] = info
    return info


def _get_lesson_path_info(db, lesson_ref: str, ctx: Dict[str, Any]) -> dict:
    lesson_ref = (lesson_ref or "").strip()
    if not lesson_ref:
        return {}
    cached = ctx.get("_lesson_path", {}).get(lesson_ref)
    if cached:
        return cached
    d = db["lesson"].find_one({"import_key": lesson_ref}, {"topic_ref": 1, "lesson_num": 1})
    if not d:
        return {}
    topic_info = _get_topic_path_info(db, (d.get("topic_ref") or "").strip(), ctx)
    if not topic_info:
        return {}
    lesson_num = _two_digit(d.get("lesson_num"))
    if not lesson_num:
        return {}
    info = {**topic_info, "lesson_num": lesson_num}
    ctx.setdefault("_lesson_path", {})[lesson_ref] = info
    return info


def _compute_asset_prefixes(
    col: str,
    doc: Dict[str, Any],
    rec: Dict[str, Any],
    db,
    ctx: Dict[str, Any],
    import_key: str,
) -> Optional[Dict[str, Any]]:
    """Compute deterministic asset_prefixes for educational entities."""
    if col == "class":
        cn = doc.get("class_name")
        if cn:
            ctx.setdefault("class", {})[import_key] = {"class_slug": _slugify_vi(cn)}
        return None

    if col == "subject":
        class_ref = str(rec.get("class_ref") or doc.get("class_ref") or "").strip()
        class_slug = _get_class_slug(db, class_ref, ctx)
        subj_slug = _slugify_vi(rec.get("subject_name") or doc.get("subject_name"))
        if not (class_slug and subj_slug):
            return None
        ctx.setdefault("_subject_path", {})[import_key] = {
            "class_slug": class_slug, "subject_slug": subj_slug,
        }
        return {
            "documents": f"documents/{class_slug}/{subj_slug}/subject",
        }

    if col == "topic":
        subject_ref = str(rec.get("subject_ref") or doc.get("subject_ref") or "").strip()
        subj_info = _get_subject_path_info(db, subject_ref, ctx)
        if not subj_info:
            return None
        n = _two_digit(rec.get("topic_num") or doc.get("topic_num"))
        if not n:
            return None
        base = f"{subj_info['class_slug']}/{subj_info['subject_slug']}"
        identifier = f"topic_{n}"
        ctx.setdefault("_topic_path", {})[import_key] = {**subj_info, "topic_num": n}
        return {
            "documents": f"documents/{base}/topic/{identifier}",
            "images": f"images/{base}/topic/{identifier}",
            "videos": f"videos/{base}/topic/{identifier}",
        }

    if col == "lesson":
        topic_ref = str(rec.get("topic_ref") or doc.get("topic_ref") or "").strip()
        topic_info = _get_topic_path_info(db, topic_ref, ctx)
        if not topic_info:
            return None
        n = _two_digit(rec.get("lesson_num") or doc.get("lesson_num"))
        if not n:
            return None
        base = f"{topic_info['class_slug']}/{topic_info['subject_slug']}"
        identifier = f"topic_{topic_info['topic_num']}-lesson_{n}"
        ctx.setdefault("_lesson_path", {})[import_key] = {**topic_info, "lesson_num": n}
        return {
            "documents": f"documents/{base}/lesson/{identifier}",
            "images": f"images/{base}/lesson/{identifier}",
            "videos": f"videos/{base}/lesson/{identifier}",
        }

    if col == "chunk":
        lesson_ref = str(rec.get("lesson_ref") or doc.get("lesson_ref") or "").strip()
        lesson_info = _get_lesson_path_info(db, lesson_ref, ctx)
        if not lesson_info:
            return None
        n = _two_digit(rec.get("chunk_num") or doc.get("chunk_num"))
        if not n:
            return None
        base = f"{lesson_info['class_slug']}/{lesson_info['subject_slug']}"
        identifier = f"topic_{lesson_info['topic_num']}-lesson_{lesson_info['lesson_num']}-chunk_{n}"
        return {
            "documents": f"documents/{base}/chunk/{identifier}",
            "images": f"images/{base}/chunk/{identifier}",
            "videos": f"videos/{base}/chunk/{identifier}",
        }

    return None


def _keyword_asset_prefixes(keyword_slug: str, mongo_id_str: str) -> Dict[str, str]:
    short_id = (mongo_id_str or "")[-6:]
    identifier = f"{keyword_slug}__{short_id}"
    return {
        "images": f"images/keyword/{identifier}",
        "videos": f"videos/keyword/{identifier}",
    }


def _ensure_keyword_related_indexes(db) -> None:
    ensure_keyword_alias_indexes(db)

    _ACTIVE = {"is_deleted": {"$ne": True}}
    try:
        db["keyword"].drop_index("import_key_1")
    except Exception:
        pass
    try:
        db["keyword"].drop_index("keyword_slug_1")
    except Exception:
        pass
    try:
        db["keyword"].drop_index("keyword_slug_1_keyword_name_1")
    except Exception:
        pass
    try:
        db["keyword"].create_index(
            "keyword_slug",
            unique=True,
            partialFilterExpression=_ACTIVE,
        )
    except Exception:
        pass

    try:
        db["keyword"].drop_index("keyword_name_1")
    except Exception:
        pass
    try:
        db["keyword"].create_index(
            "keyword_name",
            unique=True,
            partialFilterExpression=_ACTIVE,
        )
    except Exception:
        pass

    try:
        db["chunk_keyword"].drop_index("chunk_id_1_keyword_id_1")
    except Exception:
        pass
    try:
        db["chunk_keyword"].create_index(
            [("chunk_id", 1), ("keyword_id", 1)],
            unique=True,
            partialFilterExpression=_ACTIVE,
        )
    except Exception:
        pass

    try:
        db["topic_bag"].drop_index("topic_id_1")
    except Exception:
        pass
    try:
        db["topic_bag"].create_index(
            "topic_id",
            unique=True,
            partialFilterExpression=_ACTIVE,
        )
    except Exception:
        pass


def _find_or_create_keyword(db, keyword_name: str, actor: str) -> Tuple[str, str]:
    keyword_slug, existing_mongo_id = _resolve_keyword_slug(db, keyword_name)
    if existing_mongo_id:
        # backfill asset_prefixes if missing
        oid = ObjectId(existing_mongo_id) if ObjectId.is_valid(existing_mongo_id) else existing_mongo_id
        existing_doc = db["keyword"].find_one({"_id": oid}, {"asset_prefixes": 1, "keyword_slug": 1})
        if existing_doc and not existing_doc.get("asset_prefixes"):
            slug = existing_doc.get("keyword_slug") or keyword_slug
            db["keyword"].update_one(
                {"_id": oid},
                {"$set": {"asset_prefixes": _keyword_asset_prefixes(slug, existing_mongo_id)}},
            )
        return existing_mongo_id, "noop"
    enforce_canonical_name_precedence(db, keyword_name, actor)

    now = _now()
    new_id = ObjectId()
    mongo_id_str = str(new_id)
    doc = {
        "_id": new_id,
        "keyword_name": keyword_name,
        "keyword_slug": keyword_slug,
        "aliases": [],
        "asset_prefixes": _keyword_asset_prefixes(keyword_slug, mongo_id_str),
        "is_deleted": False,
        "deleted_at": None,
        "created_at": now,
        "updated_at": now,
        "created_by": actor,
        "updated_by": actor,
    }
    db["keyword"].insert_one(doc)
    return mongo_id_str, "insert"


def _upsert_chunk_keyword(db, chunk_id: str, keyword_id: str, actor: str) -> str:
    chunk_oid = ObjectId(chunk_id) if ObjectId.is_valid(chunk_id) else chunk_id
    kw_oid = ObjectId(keyword_id) if ObjectId.is_valid(keyword_id) else keyword_id

    existing = db["chunk_keyword"].find_one(
        {"chunk_id": chunk_oid, "keyword_id": kw_oid, "is_deleted": {"$ne": True}},
        {"_id": 1},
    )
    if existing:
        return "noop"

    now = _now()
    db["chunk_keyword"].insert_one({
        "chunk_id": chunk_oid,
        "keyword_id": kw_oid,
        "is_deleted": False,
        "deleted_at": None,
        "created_at": now,
        "updated_at": now,
        "created_by": actor,
        "updated_by": actor,
    })
    return "insert"


def _upsert_topic_bag(
    db, topic_id: str, topic_name: Optional[str], keyword_id: str, keyword_name: str, actor: str
) -> str:
    topic_oid = ObjectId(topic_id) if ObjectId.is_valid(topic_id) else topic_id
    kw_oid = ObjectId(keyword_id) if ObjectId.is_valid(keyword_id) else keyword_id

    now = _now()
    existing = db["topic_bag"].find_one(
        {"topic_id": topic_oid, "is_deleted": {"$ne": True}},
        {"_id": 1, "keyword_refs": 1},
    )

    if existing:
        before_refs = existing.get("keyword_refs") or []
        if any(r.get("keyword_id") == kw_oid for r in before_refs):
            return "noop"
        set_fields: Dict[str, Any] = {"updated_at": now, "updated_by": actor}
        if topic_name is not None:
            set_fields["topic_name"] = topic_name
        db["topic_bag"].update_one(
            {"_id": existing["_id"]},
            {
                "$push": {"keyword_refs": {"keyword_id": kw_oid, "keyword_name": keyword_name}},
                "$set": set_fields,
            },
        )
        updated_doc = db["topic_bag"].find_one({"_id": existing["_id"]}, {"keyword_refs": 1})
        total = len(updated_doc.get("keyword_refs") or [])
        db["topic_bag"].update_one({"_id": existing["_id"]}, {"$set": {"total_keywords": total}})
        return "update"

    db["topic_bag"].insert_one({
        "topic_id": topic_oid,
        "topic_name": topic_name,
        "keyword_refs": [{"keyword_id": kw_oid, "keyword_name": keyword_name}],
        "total_keywords": 1,
        "is_deleted": False,
        "deleted_at": None,
        "created_at": now,
        "updated_at": now,
        "created_by": actor,
        "updated_by": actor,
    })
    return "insert"


def _finalize_topic_embeddings(
    db,
    affected_topic_ids: set[str],
    sync_one: Callable[[str, Dict[str, Any]], Dict[str, Any]],
    errors: List[Dict[str, Any]],
) -> Dict[str, Any]:
    """Sync each affected topic once after the full topic_bag is populated.

    Rebuilds keyword_embedding_candidates / keyword_embedding_selected /
    keyword_embedding_text / keyword_embedding_used_fallback in Mongo and
    pushes the updated embedding to PG + Neo4j.
    """
    finalized = 0
    finalize_errors: List[Dict[str, Any]] = []

    _log.info("[import] _finalize_topic_embeddings: starting for %d topic(s)", len(affected_topic_ids))

    for topic_id in affected_topic_ids:
        try:
            topic_oid = ObjectId(topic_id) if ObjectId.is_valid(topic_id) else topic_id
            topic_doc = db["topic"].find_one({"_id": topic_oid, "is_deleted": {"$ne": True}})
            if not topic_doc:
                _log.warning(
                    "[import] finalize: topic id=%s not found or deleted — skipped",
                    topic_id,
                )
                continue

            _import_key = topic_doc.get("import_key") or ""
            _topic_name = topic_doc.get("topic_name") or ""
            _log.info(
                "[import] finalize: syncing topic id=%s import_key=%s name=%s",
                topic_id, _import_key, _topic_name,
            )

            result = sync_one("topic", topic_doc)
            if isinstance(result, dict) and result.get("ok"):
                finalized += 1
                _log.info(
                    "[import] finalize: OK topic id=%s import_key=%s",
                    topic_id, _import_key,
                )
            else:
                err = result.get("error", "no detail") if isinstance(result, dict) else "no detail"
                _log.warning(
                    "[import] finalize: FAILED topic id=%s import_key=%s error=%s",
                    topic_id, _import_key, err,
                )
                finalize_errors.append({"topic_id": topic_id, "import_key": _import_key, "error": f"sync_failed: {err}"})
        except Exception as e:
            _log.warning("[import] finalize: EXCEPTION topic id=%s: %s", topic_id, e)
            finalize_errors.append({"topic_id": topic_id, "error": str(e)})

    _log.info(
        "[import] _finalize_topic_embeddings: done — finalized=%d/%d errors=%d",
        finalized, len(affected_topic_ids), len(finalize_errors),
    )

    if finalize_errors:
        errors.extend(finalize_errors)

    return {"affected_topics": len(affected_topic_ids), "finalized_topics": finalized, "topic_finalize_errors": finalize_errors}


def _import_keyword_rows(
    db,
    rows: List[Dict[str, Any]],
    actor: str,
    *,
    progress_callback: Optional[Callable[[Dict[str, Any]], None]] = None,
    progress_state: Optional[Dict[str, Any]] = None,
    sync_one: Optional[Callable[[str, Dict[str, Any]], Dict[str, Any]]] = None,
    minio_client=None,
    minio_bucket: Optional[str] = None,
    minio_seen: Optional[Set[str]] = None,
    minio_errors: Optional[List[Dict[str, Any]]] = None,
) -> Dict[str, Any]:

    inserted = reused = synced = 0
    errors: List[Dict[str, Any]] = []
    new_keywords: Dict[str, str] = {}
    reused_keyword_ids: set[str] = set()
    affected_topic_ids: set[str] = set()

    for rec in rows:
        rowno = rec.pop("_row", None)
        try:
            chunk_ref = str(rec.get("chunk_ref") or "").strip()
            keyword_name = str(rec.get("keyword_name") or "").strip()

            if not chunk_ref:
                raise ValueError("missing chunk_ref")
            if not keyword_name:
                raise ValueError("missing keyword_name")

            chunk_doc = db["chunk"].find_one(
                {"import_key": chunk_ref, "is_deleted": {"$ne": True}},
                {"_id": 1, "lesson_id": 1},
            )
            if not chunk_doc:
                raise ValueError(f"chunk with import_key='{chunk_ref}' not found")
            chunk_id = str(chunk_doc["_id"])

            lesson_id = str(chunk_doc.get("lesson_id") or "").strip()
            if not lesson_id:
                raise ValueError(f"chunk '{chunk_ref}' has no lesson_id")

            lesson_doc = None
            if ObjectId.is_valid(lesson_id):
                lesson_doc = db["lesson"].find_one(
                    {"_id": ObjectId(lesson_id), "is_deleted": {"$ne": True}},
                    {"topic_id": 1},
                )
            if not lesson_doc:
                lesson_doc = db["lesson"].find_one(
                    {"_id": lesson_id, "is_deleted": {"$ne": True}},
                    {"topic_id": 1},
                )
            if not lesson_doc:
                raise ValueError(f"lesson '{lesson_id}' not found")
            topic_id = str(lesson_doc.get("topic_id") or "").strip()
            if not topic_id:
                raise ValueError(f"lesson '{lesson_id}' has no topic_id")

            topic_doc = None
            if ObjectId.is_valid(topic_id):
                topic_doc = db["topic"].find_one(
                    {"_id": ObjectId(topic_id), "is_deleted": {"$ne": True}},
                    {"topic_name": 1},
                )
            if not topic_doc:
                topic_doc = db["topic"].find_one(
                    {"_id": topic_id, "is_deleted": {"$ne": True}},
                    {"topic_name": 1},
                )
            if not topic_doc:
                raise ValueError(f"topic '{topic_id}' not found")
            topic_name: Optional[str] = topic_doc.get("topic_name") or None

            keyword_id, kw_op = _find_or_create_keyword(db, keyword_name, actor)

            if minio_client and minio_bucket:
                _kw_oid = ObjectId(keyword_id) if ObjectId.is_valid(keyword_id) else keyword_id
                _kw_doc = db["keyword"].find_one({"_id": _kw_oid}, {"asset_prefixes": 1})
                if _kw_doc and _kw_doc.get("asset_prefixes"):
                    ensure_asset_prefix_markers(
                        minio_client, minio_bucket, _kw_doc["asset_prefixes"],
                        seen=minio_seen, errors=minio_errors,
                    )

            if kw_op == "insert":
                new_keywords.setdefault(keyword_id, keyword_name)
                inserted += 1
            else:
                reused_keyword_ids.add(keyword_id)
                reused += 1

            _upsert_chunk_keyword(db, chunk_id, keyword_id, actor)
            _upsert_topic_bag(db, topic_id, topic_name, keyword_id, keyword_name, actor)

            # Always mark topic for finalization regardless of bag_op.
            # Topics with a pre-existing topic_bag (noop) still need their
            # keyword_embedding_* fields rebuilt when they were empty (e.g. on re-import).
            if topic_id not in affected_topic_ids:
                _log.debug(
                    "[import] topic marked for finalization: id=%s name=%s",
                    topic_id, topic_name or "",
                )
            affected_topic_ids.add(topic_id)

            if sync_one is not None:
                try:
                    _ck_chunk_oid = ObjectId(chunk_id) if ObjectId.is_valid(chunk_id) else chunk_id
                    _ck_kw_oid = ObjectId(keyword_id) if ObjectId.is_valid(keyword_id) else keyword_id
                    ck_doc = db["chunk_keyword"].find_one(
                        {"chunk_id": _ck_chunk_oid, "keyword_id": _ck_kw_oid, "is_deleted": {"$ne": True}}
                    )
                    if ck_doc:
                        sync_result = sync_one("chunk_keyword", ck_doc)
                        if isinstance(sync_result, dict) and sync_result.get("ok"):
                            synced += 1
                        elif isinstance(sync_result, dict):
                            errors.append({"row": rowno, "error": f"sync_failed: {sync_result.get('error') or 'no detail'}", "collection": "chunk_keyword"})
                except Exception as sync_e:
                    errors.append({"row": rowno, "error": f"sync_exception: {sync_e}", "collection": "chunk_keyword"})

        except Exception as e:
            errors.append({"row": rowno, "error": str(e), "collection": "keyword"})
        finally:
            if progress_callback is not None and progress_state is not None:
                progress_state["processed_rows"] += 1
                _pr = progress_state["processed_rows"]
                _tot = progress_state["total_rows"]
                progress_callback({
                    "current_collection": "keyword",
                    "processed_rows": _pr,
                    "total_rows": _tot,
                    "progress": min(int(_pr * 100 / _tot), 99) if _tot > 0 else 99,
                    "message": "Đang import keyword...",
                })

    alias_skipped = len(reused_keyword_ids - new_keywords.keys())

    phase2_slots = len(rows)
    batch_result = {
        "processed_keywords": 0,
        "inserted_aliases": 0,
        "stopped_due_to_quota": False,
        "remaining_keywords": [],
    }

    if new_keywords:
        from app.services.keyword_alias_service import refresh_keyword_aliases_batch

        try:
            batch_result = refresh_keyword_aliases_batch(
                db=db,
                keyword_id_name_pairs=list(new_keywords.items()),
                actor=actor,
                batch_size=5,
                screen_batch_size=25,
                batch_sleep=6.0,
                max_wait_seconds=3600,
                progress_callback=progress_callback,
                progress_state=progress_state,
                phase2_total_slots=phase2_slots,
            )
        except Exception as batch_e:
            _log.warning("[import] alias batch failed: %s", batch_e)

    if progress_callback is not None and progress_state is not None and phase2_slots > 0:
        already_emitted = progress_state.get("_alias_slots_emitted", 0)
        remaining = phase2_slots - already_emitted
        if remaining > 0:
            progress_state["processed_rows"] = progress_state.get("processed_rows", 0) + remaining
            _pr = progress_state["processed_rows"]
            _tot = progress_state.get("total_rows", 0)
            progress_callback({
                "current_collection": "keyword",
                "processed_rows": _pr,
                "total_rows": _tot,
                "progress": min(int(_pr * 100 / _tot), 99) if _tot > 0 else 99,
                "message": "Hoàn tất import keyword.",
            })
        else:
            _pr = progress_state.get("processed_rows", 0)
            _tot = progress_state.get("total_rows", 0)
            progress_callback({
                "current_collection": "keyword",
                "processed_rows": _pr,
                "total_rows": _tot,
                "progress": min(int(_pr * 100 / _tot), 99) if _tot > 0 else 99,
                "message": "Hoàn tất import keyword.",
            })

    topic_finalize_summary: Dict[str, Any] = {"affected_topics": 0, "finalized_topics": 0, "topic_finalize_errors": []}
    if sync_one is not None and affected_topic_ids:
        _log.info("[import] finalizing topic embeddings for %d affected topic(s)", len(affected_topic_ids))
        topic_finalize_summary = _finalize_topic_embeddings(db, affected_topic_ids, sync_one, errors)

    return {
        "rows": len(rows),
        "inserted": inserted,
        "reused": reused,
        "synced": synced,
        "errors": errors[:50],
        "alias_processed_keywords": batch_result["processed_keywords"],
        "alias_inserted": batch_result["inserted_aliases"],
        "alias_deleted": 0,
        "alias_skipped": alias_skipped,
        "alias_errors": [],
        "alias_stopped_due_to_quota": batch_result["stopped_due_to_quota"],
        "alias_remaining_keywords": batch_result["remaining_keywords"],
        **topic_finalize_summary,
    }


def import_excel_to_mongo(
    db,
    xlsx_path: str,
    *,
    actor: str,
    sync_one: Optional[Callable[[str, Dict[str, Any]], Dict[str, Any]]] = None,
    only_cols: Optional[List[str]] = None,
    progress_callback: Optional[Callable[[Dict[str, Any]], None]] = None,
) -> Dict[str, Any]:
    wb = load_workbook(xlsx_path, data_only=True)

    cols = only_cols or IMPORT_ORDER
    all_cols = set(IMPORT_ORDER) | set(cols)
    id_map: Dict[str, Dict[str, str]] = {c: {} for c in all_cols}

    ctx: Dict[str, Any] = {"class": {}, "_subject_path": {}, "_topic_path": {}, "_lesson_path": {}}

    minio_client, minio_bucket = _get_import_minio()
    minio_seen: Set[str] = set()
    minio_errors: List[Dict[str, Any]] = []
    if minio_client:
        ensure_root_folders(minio_client, minio_bucket, errors=minio_errors)

    report = {"file": xlsx_path, "collections": {}, "errors": []}

    rows_by_col: Dict[str, List[Dict[str, Any]]] = {col: _read_sheet_rows(wb, col) for col in cols}
    total_rows: int = sum(
        len(r) * 2 if col == "keyword" else len(r)
        for col, r in rows_by_col.items()
    )
    processed_rows: int = 0

    for col in cols:
        rows = rows_by_col[col]
        if not rows:
            report["collections"][col] = {"rows": 0, "inserted": 0, "updated": 0, "synced": 0, "skipped": True}
            continue

        if col == "keyword":
            _ensure_keyword_related_indexes(db)
            _progress_state = (
                {"processed_rows": processed_rows, "total_rows": total_rows}
                if progress_callback is not None else None
            )
            report["collections"]["keyword"] = _import_keyword_rows(
                db, rows, actor,
                progress_callback=progress_callback,
                progress_state=_progress_state,
                sync_one=sync_one,
                minio_client=minio_client,
                minio_bucket=minio_bucket,
                minio_seen=minio_seen,
                minio_errors=minio_errors,
            )
            if _progress_state is not None:
                processed_rows = _progress_state["processed_rows"]
            continue

        _ensure_import_index(db, col)

        inserted = updated = synced = 0
        errors = []

        for rec in rows:
            rowno = rec.pop("_row", None)
            try:
                import_key = str(rec.get("import_key") or "").strip()
                if not import_key:
                    raise ValueError("missing import_key")

                doc: Dict[str, Any] = {}
                for k, v in rec.items():
                    if k is None:
                        continue
                    key = str(k).strip()
                    if key == "" or key == "import_key":
                        continue

                    if key in JSON_FIELDS:
                        doc[key] = _try_parse_json(v)
                    else:
                        doc[key] = _cell_to_value(v)

                if col in REF_MAP:
                    ref_col, target_field, parent_col = REF_MAP[col]
                    ref_key = str(rec.get(ref_col) or "").strip()
                    if ref_key:
                        parent_id = id_map[parent_col].get(ref_key)
                        if not parent_id:
                            parent_doc = db[parent_col].find_one({"import_key": ref_key}, {"_id": 1})
                            if parent_doc:
                                parent_id = str(parent_doc["_id"])
                                id_map[parent_col][ref_key] = parent_id
                        if not parent_id:
                            raise ValueError(
                                f"cannot resolve {ref_col}='{ref_key}' (parent '{parent_col}' not imported yet)"
                            )
                        doc[target_field] = ObjectId(parent_id) if ObjectId.is_valid(parent_id) else parent_id

                asset_prefixes = _compute_asset_prefixes(col, doc, rec, db, ctx, import_key)
                if asset_prefixes is not None:
                    doc["asset_prefixes"] = asset_prefixes

                doc["import_key"] = import_key

                mongo_id, op = _upsert_by_import_key(db, col, import_key, doc, actor=actor)

                id_map[col][import_key] = mongo_id

                if asset_prefixes and minio_client:
                    ensure_asset_prefix_markers(
                        minio_client, minio_bucket, asset_prefixes,
                        seen=minio_seen, errors=minio_errors,
                    )

                if op == "insert":
                    inserted += 1
                elif op == "update":
                    updated += 1

                if sync_one and op != "noop":
                    full = db[col].find_one({"import_key": import_key})
                    if full:
                        try:
                            sync_result = sync_one(col, full)
                            if isinstance(sync_result, dict) and sync_result.get("ok"):
                                synced += 1
                            elif isinstance(sync_result, dict):
                                errors.append({"row": rowno, "error": f"sync_failed: {sync_result.get('error') or 'no detail'}", "collection": col})
                        except Exception as sync_e:
                            errors.append({"row": rowno, "error": f"sync_exception: {sync_e}", "collection": col})

            except Exception as e:
                errors.append({"row": rowno, "error": str(e), "collection": col})
            finally:
                if progress_callback is not None:
                    processed_rows += 1
                    progress_callback({
                        "current_collection": col,
                        "processed_rows": processed_rows,
                        "total_rows": total_rows,
                        "progress": min(int(processed_rows * 100 / total_rows), 99) if total_rows > 0 else 99,
                        "message": f"Đang xử lý {col}...",
                    })

        report["collections"][col] = {
            "rows": len(rows),
            "inserted": inserted,
            "updated": updated,
            "synced": synced,
            "errors": errors[:50],
        }

    if minio_errors:
        report["minio_errors"] = minio_errors[:50]

    return report